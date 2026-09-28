import asyncio
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from lxml import etree

from app.enterprise_capabilities.docx_review.annotations import annotate_docx
from app.enterprise_capabilities.docx_review.source import page_blocks, qn, read_document_xml, review_blocks, source_hash
from app.enterprise_capabilities.docx_review import service as review_service
from app.enterprise_capabilities.runtime.contracts import CapabilityExecutionContext


def _sample() -> bytes:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Service fee: ").bold = True
    paragraph.add_run("100 units, payable on delivery.")
    document.add_paragraph("Service fee: 100 units, payable on delivery.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Acceptance occurs after written confirmation."
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_annotations_bind_exact_source_ranges_without_rebuilding_docx():
    source = _sample()
    blocks = review_blocks(read_document_xml(source))
    assert len(blocks) == 3
    findings = [
        {"source_block_id": blocks[0].id, "source_quote": "fee: 100 units", "criterion_ref": "C-1",
         "criterion_text": "Pay after acceptance", "finding": "Fee should be confirmed", "suggested_revision": "Specify the amount"},
        {"source_block_id": blocks[2].id, "source_quote": "written confirmation", "criterion_ref": "C-2",
         "criterion_text": "Acceptance requires a deadline", "finding": "Acceptance may be delayed",
         "suggested_revision": "Set a written acceptance deadline"},
        {"source_block_id": blocks[0].id, "source_quote": "no such words", "criterion_ref": "C-3",
         "criterion_text": "Use actual source text", "finding": "Not grounded", "suggested_revision": "Correct the quote"},
    ]
    output, accepted, rejected = annotate_docx(source, findings)
    assert output is not None
    assert len(accepted) == 2
    assert rejected[0]["reason"] == "quote_not_found"
    assert [(block.id, block.text) for block in review_blocks(read_document_xml(output))] == [
        (block.id, block.text) for block in blocks
    ]
    assert source_hash(output) != source_hash(source)
    with ZipFile(BytesIO(source)) as before, ZipFile(BytesIO(output)) as after:
        assert before.read("word/styles.xml") == after.read("word/styles.xml")
        comments = etree.fromstring(after.read("word/comments.xml"))
        assert len(comments.findall(qn("comment"))) == 2
        first_comment = comments.findall(qn("comment"))[0]
        assert len(first_comment.findall(qn("p"))) == 1
        assert len(first_comment.findall(f".//{qn('br')}")) == 3
        lines = [node.text for node in first_comment.iter(qn("t"))]
        assert lines == [
            "问题：Fee should be confirmed", "修改建议：Specify the amount",
            "审核依据：Pay after acceptance", "规则编号：C-1",
        ]
        assert len(read_document_xml(output).findall(f".//{qn('commentRangeStart')}")) == 2
        assert b"Not grounded" not in after.read("word/comments.xml")
    assert len(Document(BytesIO(output)).tables) == 1


def test_no_unverified_comments_are_delivered():
    source = _sample()
    output, accepted, rejected = annotate_docx(source, [
        {"source_block_id": "p999999", "source_quote": "Service fee", "finding": "Wrong block"},
    ])
    assert output is None
    assert accepted == []
    assert rejected[0]["reason"] == "source_block_not_found"


def test_rule_identifier_alone_cannot_become_a_user_comment():
    source = _sample()
    output, accepted, rejected = annotate_docx(source, [{
        "source_block_id": "p000001", "source_quote": "fee: 100 units",
        "criterion_ref": "R13", "finding": "Check the fee",
    }])
    assert output is None and accepted == []
    assert rejected[0]["reason"] == "incomplete_finding"


def test_source_pages_cover_every_block_in_original_order():
    document = Document()
    for index in range(5):
        document.add_paragraph(f"Clause {index}: " + "x" * 600)
    stream = BytesIO()
    document.save(stream)
    blocks = review_blocks(read_document_xml(stream.getvalue()))
    first = page_blocks(blocks, offset=0, max_chars=1000)
    assert len(first["blocks"]) == 1 and first["has_more"]
    ids = [first["blocks"][0]["id"]]
    offset = first["next_offset"]
    while offset is not None:
        page = page_blocks(blocks, offset=offset, max_chars=1000)
        ids.extend(item["id"] for item in page["blocks"])
        offset = page["next_offset"]
    assert ids == [block.id for block in blocks]


def test_ambiguous_quote_does_not_become_a_comment():
    document = Document()
    document.add_paragraph("same phrase; same phrase")
    stream = BytesIO()
    document.save(stream)
    output, accepted, rejected = annotate_docx(stream.getvalue(), [
        {"source_block_id": "p000001", "source_quote": "same phrase", "criterion_ref": "C-1",
         "criterion_text": "Unique reference", "finding": "Ambiguous", "suggested_revision": "Clarify phrase"},
    ])
    assert output is None
    assert accepted == []
    assert rejected[0]["reason"] == "ambiguous_quote"


def test_cached_page_break_inside_run_keeps_exact_anchor_and_single_marker():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("non-exclusive ")
    marked = paragraph.add_run("transferable license")
    marked._r.insert(1, OxmlElement("w:lastRenderedPageBreak"))
    stream = BytesIO()
    document.save(stream)
    source = stream.getvalue()
    output, accepted, rejected = annotate_docx(source, [{
        "source_block_id": "p000001", "source_quote": "exclusive transferable",
        "criterion_text": "A defined license is required", "finding": "Scope is unclear",
        "suggested_revision": "Define scope and duration",
    }])
    assert output is not None and len(accepted) == 1 and rejected == []
    root = read_document_xml(output)
    assert review_blocks(root)[0].text == review_blocks(read_document_xml(source))[0].text
    assert len(list(root.iter(qn("lastRenderedPageBreak")))) == 1


def test_review_tools_use_the_same_authorized_original(monkeypatch):
    source = _sample()
    uploads = []
    monkeypatch.setattr(review_service, "read_owned_artifact", lambda artifact, *, context: (
        artifact, source, "sample.docx",
    ))
    monkeypatch.setattr(review_service, "upload_derived_artifact", lambda content, *, context, filename, content_type: (
        uploads.append((content, filename, content_type)) or {"object_path": "owner/reviewed.docx", "filename": filename}
    ))
    context = CapabilityExecutionContext(
        tenant_id="tenant", user_id="owner", conversation_id="conversation",
        kernel_session_id="kernel", profile_version="profile", action_id="action",
        turn_context={"documents": [{"object_path": "owner/sample.docx", "filename": "sample.docx"}]},
    )
    artifact = {"object_path": "owner/sample.docx", "filename": "sample.docx"}
    read = asyncio.run(review_service.document_review_source({"artifact": artifact}, context))
    assert read["success"] and not read["has_more"]
    assert read["source_sha256"] == source_hash(source)
    result = asyncio.run(review_service.document_annotate({
        "artifact": artifact, "source_sha256": read["source_sha256"],
        "findings": [{"source_block_id": read["blocks"][0]["id"], "source_quote": "fee: 100 units",
                      "criterion_ref": "C-1", "criterion_text": "Pay after acceptance",
                      "finding": "Check fee", "suggested_revision": "Specify the amount"}],
    }, context))
    assert result["success"] and result["accepted_count"] == 1
    assert uploads[0][1] == "sample_AI审阅.docx"
    assert review_blocks(read_document_xml(uploads[0][0]))[0].text == read["blocks"][0]["text"]


def test_review_rejects_an_owned_but_not_currently_uploaded_docx(monkeypatch):
    def unexpected_read(*args, **kwargs):
        raise AssertionError("storage must not be read before turn attachment authorization")

    monkeypatch.setattr(review_service, "read_owned_artifact", unexpected_read)
    context = CapabilityExecutionContext(
        tenant_id="tenant", user_id="owner", conversation_id="conversation",
        kernel_session_id="kernel", profile_version="profile", action_id="action",
        turn_context={"documents": [{"object_path": "owner/current.docx"}]},
    )
    try:
        asyncio.run(review_service.document_review_source({
            "artifact": {"object_path": "owner/older.docx", "filename": "older.docx"},
        }, context))
    except PermissionError as exc:
        assert "current turn" in str(exc)
    else:
        raise AssertionError("older document was allowed")
